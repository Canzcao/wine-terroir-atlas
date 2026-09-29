# coding: utf-8
"""把缺数据清单做成 Excel（openpyxl）：可筛选、可勾进度、带汇总与话术。

产出：outputs/collect/_data-gaps-2026-09-21/酒庄缺数据清单.xlsx
Sheet1 缺失清单（筛选+冻结+分组配色）
Sheet2 按国家汇总
Sheet3 索要话术模板
Sheet4 使用说明
"""
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SEED = ROOT / 'data' / 'catalog-seed.json'
OUT = ROOT / 'outputs' / 'collect' / '_data-gaps-2026-09-21'
ORDER = ['邮箱', '联系人', '电话', '坐标', '官网', '地址', '酒款', '图片', '本地语资料']
FILLS = {'A 可直接联系': 'E8F5E9', 'B 需先找联系人': 'FFF3E0', 'C 仅缺少量': 'E3F2FD'}
HEAD = PatternFill('solid', fgColor='4A148C')
HEADF = Font(color='FFFFFF', bold=True, size=11)
THIN = Border(*[Side(style='thin', color='DDDDDD')] * 4)


def main():
    seed = json.loads(SEED.read_text())
    recs = seed if isinstance(seed, list) else seed.get('records', [])
    rname = {r['id']: (r.get('data') or {}).get('name') or r.get('name')
             for r in recs if r.get('kind') == 'region'}
    wines = Counter()
    for r in recs:
        if r.get('kind') == 'wine':
            w = (r.get('data') or {}).get('wineryId')
            if w:
                wines[w] += 1
    rows = []
    for r in recs:
        if r.get('kind') != 'winery':
            continue
        d = r.get('data') or {}
        rid = d.get('regionId')
        miss = []
        if not d.get('email'):
            miss.append('邮箱')
        if not d.get('contactPerson'):
            miss.append('联系人')
        if not d.get('phone'):
            miss.append('电话')
        if d.get('lat') is None:
            miss.append('坐标')
        if not d.get('website'):
            miss.append('官网')
        if not d.get('address'):
            miss.append('地址')
        if wines.get(r['id'], 0) == 0:
            miss.append('酒款')
        if not (d.get('images') or []):
            miss.append('图片')
        if not (d.get('localizations') or {}):
            miss.append('本地语资料')
        miss.sort(key=lambda x: ORDER.index(x))
        reach = bool(d.get('website') or d.get('email'))
        grp = ('C 仅缺少量' if 1 <= len(miss) <= 2 else 'A 可直接联系') if reach else 'B 需先找联系人'
        rows.append([grp, d.get('country') or '', rname.get(rid, rid or ''), r.get('name'),
                     d.get('website') or '', d.get('email') or '', d.get('phone') or '',
                     d.get('address') or '', '有' if d.get('lat') is not None else '无',
                     wines.get(r['id'], 0), len(d.get('images') or []),
                     '、'.join(miss), len(miss), '、'.join(miss[:3]), d.get('sourceURL') or '', r['id'], '', ''])
    rows.sort(key=lambda x: (x[0], x[1], -x[12], x[3]))

    wb = Workbook()
    ws = wb.active
    ws.title = '缺失清单'
    cols = ['组', '国家', '产区', '酒庄（原语种名）', '官网', '邮箱', '电话', '地址', '有坐标',
            '酒款数', '图片数', '缺失项', '缺几项', '最该先问的3件', '来源页', '记录ID', '联系状态', '回复/备注']
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD; cell.font = HEADF; cell.alignment = Alignment(vertical='center')
    for row in rows:
        ws.append(row)
        f = PatternFill('solid', fgColor=FILLS.get(row[0], 'FFFFFF'))
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=ws.max_row, column=c)
            cell.fill = f; cell.border = THIN
            cell.alignment = Alignment(vertical='top', wrap_text=(c in (12, 14)))
    widths = [14, 10, 16, 34, 30, 26, 18, 30, 8, 8, 8, 30, 8, 24, 34, 34, 14, 20]
    for i, wd in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = 'A1:R%d' % ws.max_row

    w2 = wb.create_sheet('按国家汇总')
    w2.append(['国家', '酒庄数', '可联系上', '缺坐标', '缺官网', '缺邮箱', '缺地址', '缺酒款', '缺图片', '缺本地语资料'])
    for c in range(1, 11):
        w2.cell(row=1, column=c).fill = HEAD
        w2.cell(row=1, column=c).font = HEADF
    agg = defaultdict(lambda: Counter())
    tot = Counter()
    for r in rows:
        a = agg[r[1]]; a['total'] += 1; tot['total'] += 1
        if r[4] or r[5]:
            a['reach'] += 1
        for m in r[11].split('、'):
            if m:
                a[m] += 1
    for c, a in sorted(agg.items(), key=lambda x: -x[1]['total']):
        w2.append([c, a['total'], a['reach'], a['坐标'], a['官网'], a['邮箱'], a['地址'], a['酒款'], a['图片'], a['本地语资料']])
    for i, wd in enumerate([12, 10, 10, 9, 9, 9, 9, 9, 9, 13], start=1):
        w2.column_dimensions[get_column_letter(i)].width = wd
    w2.freeze_panes = 'A2'

    w3 = wb.create_sheet('索要话术模板')
    w3.append(['语言', '模板（把 {region} 换成产区名、{items} 换成缺失项）'])
    w3.cell(row=1, column=1).fill = HEAD; w3.cell(row=1, column=1).font = HEADF
    w3.cell(row=1, column=2).fill = HEAD; w3.cell(row=1, column=2).font = HEADF
    for lang, t in json.loads((OUT / '缺口汇总.json').read_text()).get('templates', {}).items():
        w3.append([lang, t])
    w3.column_dimensions['A'].width = 14
    w3.column_dimensions['B'].width = 110

    w4 = wb.create_sheet('使用说明')
    for line in [
        ['酒庄缺数据清单 — 使用说明'],
        [''],
        ['1. 「缺失清单」页已开筛选与冻结首行：按「组」筛 A/B/C，按「国家」筛一国之内的酒庄。'],
        ['2. 颜色含义：绿色=A 可直接联系（有官网或邮箱）；橙色=B 需先找联系人（连官网/邮箱都没有）；蓝色=C 仅缺 1–2 项，顺手补齐。'],
        ['3. 最后两列「联系状态」「回复/备注」留给你自己填，例如：已发邮件 9/22、已回-待补坐标、无回应-10/05 再发。'],
        ['4. 「最该先问的3件」列已按优先级排序：邮箱 > 联系人 > 电话 > 坐标 > 官网 > 地址 > 酒款 > 图片 > 本地语资料。'],
        ['5. 话术模板见「索要话术模板」页，已按酒庄所在国语言准备（中/英/法/西/意/德/葡）。'],
        ['6. 边界：只索要资料，不承诺商业回报；图片说明"标注来源、侵权可联系删除"。'],
        ['7. 收到回复后按 采集规格-网站字段对照.md 录入，并记 sourceURL 与 checkedDate。'],
    ]:
        w4.append(line)
    w4.column_dimensions['A'].width = 120
    w4['A1'].font = Font(bold=True, size=14)

    path = OUT / '酒庄缺数据清单.xlsx'
    wb.save(path)
    print('已生成 %s（%d 行酒庄）' % (path, len(rows)))


if __name__ == '__main__':
    main()
